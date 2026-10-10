from __future__ import annotations

import io
import zipfile

from docx import Document


def make_text_pdf(pages: list[str]) -> bytes:
    objects: dict[int, bytes] = {}
    font_id = 3
    objects[1] = b"<< /Type /Catalog /Pages 2 0 R >>"
    kids: list[str] = []
    page_entries: list[tuple[int, int, str]] = []
    next_id = 4
    for text in pages:
        page_id = next_id
        content_id = next_id + 1
        next_id += 2
        page_entries.append((page_id, content_id, text))
        kids.append(f"{page_id} 0 R")
    objects[2] = f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(pages)} >>".encode()
    objects[font_id] = b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"
    for page_id, content_id, text in page_entries:
        objects[page_id] = (
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 {font_id} 0 R >> >> /Contents {content_id} 0 R >>"
        ).encode()
        escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        stream = f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET".encode("latin-1")
        objects[content_id] = b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream"
    output = bytearray(b"%PDF-1.4\n")
    offsets: dict[int, int] = {}
    for object_id in sorted(objects):
        offsets[object_id] = len(output)
        output += f"{object_id} 0 obj\n".encode() + objects[object_id] + b"\nendobj\n"
    xref_offset = len(output)
    max_id = max(objects)
    output += f"xref\n0 {max_id + 1}\n".encode()
    output += b"0000000000 65535 f \n"
    for object_id in range(1, max_id + 1):
        output += f"{offsets[object_id]:010d} 00000 n \n".encode()
    output += f"trailer\n<< /Size {max_id + 1} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF\n".encode()
    return bytes(output)


def make_blank_pdf(page_count: int = 1) -> bytes:
    from pypdf import PdfWriter

    writer = PdfWriter()
    for _ in range(page_count):
        writer.add_blank_page(width=612, height=792)
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def make_encrypted_pdf() -> bytes:
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    writer.encrypt("secret-pass")
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def make_docx(
    *,
    heading: str | None = None,
    paragraphs: list[str] | None = None,
    tables: list[list[list[str]]] | None = None,
) -> bytes:
    document = Document()
    if heading is not None:
        document.add_heading(heading, level=1)
    for paragraph in paragraphs or []:
        document.add_paragraph(paragraph)
    for table_rows in tables or []:
        table = document.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for row_index, row in enumerate(table_rows):
            for column_index, value in enumerate(row):
                table.cell(row_index, column_index).text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_zip_bomb_docx() -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", "<document>" + "A" * 2_000_000 + "</document>")
    return buffer.getvalue()


def make_macro_docx() -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("word/document.xml", "<w:document xmlns:w='urn:w'/>")
        archive.writestr("word/vbaProject.bin", b"macro-bytes")
    return buffer.getvalue()


def make_external_entity_docx() -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr(
            "word/document.xml",
            "<!DOCTYPE foo [<!ENTITY xxe SYSTEM 'file:///etc/passwd'>]><doc>&xxe;</doc>",
        )
    return buffer.getvalue()


def make_traversal_docx() -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("word/document.xml", "<w:document xmlns:w='urn:w'/>")
        archive.writestr("../evil.xml", "<evil/>")
    return buffer.getvalue()
